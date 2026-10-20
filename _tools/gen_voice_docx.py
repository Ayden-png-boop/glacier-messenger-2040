# -*- coding: utf-8 -*-
"""把配音脚本 CSV 转成排版好的 Word 文档（A4 横向，中英对照表格）。"""
import csv, io, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = r'C:/Users/-/Desktop/冰川信使海外版/output/voice-script'
CSV = os.path.join(BASE, '配音脚本-中英对照.csv')
OUT = os.path.join(BASE, '冰川信使2040-配音脚本-中英对照.docx')

rows = list(csv.DictReader(io.open(CSV, encoding='utf-8-sig')))

doc = Document()

# 页面：A4 横向
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(1.6)
sec.top_margin = sec.bottom_margin = Cm(1.8)

# 默认字体
st = doc.styles['Normal']
st.font.name = 'Microsoft YaHei'
st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


def _cell_font(cell, size=9, bold=False, color=None):
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.name = 'Microsoft YaHei'
            r.font.size = Pt(size)
            r.font.bold = bold
            if color:
                r.font.color.rgb = color
            r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


def add_table(headers, data, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        _cell_font(c, size=9, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))
    for row in data:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            _cell_font(cells[i], size=9)
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Cm(w)
    return t


# ── 封面标题 ──
h = doc.add_paragraph()
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h.add_run('《冰川信使：斑头雁的2040》\n配音脚本（中英对照）')
r.font.size = Pt(22)
r.font.bold = True
r.font.name = 'Microsoft YaHei'
r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = sub.add_run('Glacier Messenger: The Bar-headed Goose of 2040 — Bilingual Voice-over Script')
rs.font.size = Pt(11)
rs.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

# ── 使用说明 ──
note = doc.add_paragraph()
nr = note.add_run(
    '使用说明：\n'
    '1. 需配音：「旁白」「角色台词」「选项」「后果提示」「条件提示」「章节/结局标题」。\n'
    '   「画面提示」是镜头调度说明，不配音，仅供参考；「UI 文案」是界面文字，游戏里已生成过音频，通常无需重录。\n'
    '2. 时长按中文 250 字/分钟、英文 140 词/分钟估算，仅供排期参考，实际以录音为准。\n'
    '3. 每条编号（VO-xxx）与游戏内剧情数据一一对应，配音完成后可按编号回贴音频文件。\n'
    '4. 中英文条目已逐句对齐，同一编号即同一语义，可分别录制中文轨与英文轨。'
)
nr.font.size = Pt(10)
nr.font.name = 'Microsoft YaHei'
nr._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

# ── 汇总 ──
from collections import Counter
cnt = Counter(r['类型'] for r in rows)
total_zh = sum(float(r['中文时长(秒)']) for r in rows if r['类型'] != '画面提示（不配音）')
total_en = sum(float(r['EN时长(秒)']) for r in rows if r['类型'] != '画面提示（不配音）')

doc.add_paragraph()
hs = doc.add_paragraph()
hr = hs.add_run('一、汇总')
hr.font.size = Pt(14)
hr.font.bold = True

add_table(
    ['项目', '数量', '项目', '数量'],
    [
        ['总条目', len(rows), '中文净配音时长（估）', f'{round(total_zh)} 秒 ≈ {int(total_zh//60)} 分 {round(total_zh%60)} 秒'],
        ['旁白', cnt.get('旁白', 0), '英文净配音时长（估）', f'{round(total_en)} 秒 ≈ {int(total_en//60)} 分 {round(total_en%60)} 秒'],
        ['角色台词', cnt.get('角色台词', 0), '章节标题', cnt.get('章节标题', 0)],
        ['选项', cnt.get('选项', 0), '结局标题', cnt.get('结局标题', 0)],
        ['后果提示', cnt.get('后果提示', 0), '画面提示（不配音）', cnt.get('画面提示（不配音）', 0)],
    ],
    widths=[4.5, 3.0, 6.0, 6.0],
)

# ── 分场景正文 ──
doc.add_page_break()
hs = doc.add_paragraph()
hr = hs.add_run('二、分场景脚本')
hr.font.size = Pt(14)
hr.font.bold = True

cur = None
for r in rows:
    if r['场景'] != cur:
        cur = r['场景']
        p = doc.add_paragraph()
        pr = p.add_run(cur)
        pr.font.size = Pt(12)
        pr.font.bold = True
        pr.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        add_table(
            ['编号', '类型', '说话者', '中文', 'English', '中/英时长'],
            [],
            widths=[1.9, 3.0, 2.2, 8.4, 8.4, 2.2],
        )
        tbl = doc.tables[-1]
    tbl = doc.tables[-1]
    cells = tbl.add_row().cells
    vals = [
        r['编号'], r['类型'], r['说话者'] or '—',
        r['中文'], r['English'], f"{r['中文时长(秒)']}s / {r['EN时长(秒)']}s",
    ]
    for i, v in enumerate(vals):
        cells[i].text = str(v)
        _cell_font(cells[i], size=8.5)
    for i, w in enumerate([1.9, 3.0, 2.2, 8.4, 8.4, 2.2]):
        cells[i].width = Cm(w)

doc.save(OUT)
print('saved:', OUT)
print('rows:', len(rows))
